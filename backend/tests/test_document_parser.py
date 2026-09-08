"""Tests for the document text parsers (Phase 4, §7)."""

import io

from app.services import document_parser
from app.services.document_parser import parse_document


def test_txt_parse(sample_txt):
    parsed = parse_document(sample_txt.encode("utf-8"), file_type="txt", filename="a.txt")
    assert parsed.file_type == "txt"
    assert parsed.full_text.startswith("AgentFlow AI is a multi-tenant")
    assert parsed.pages and parsed.pages[0].text


def test_markdown_parse():
    md = "# Title\n\nBody paragraph with **formatting**.\n\n## Section\n\nMore text.\n" * 10
    parsed = parse_document(md.encode("utf-8"), file_type="md", filename="b.md")
    assert parsed.file_type == "md"
    assert "Title" in parsed.full_text
    assert "formatting" in parsed.full_text


def test_pdf_parse_preserves_pages(sample_pdf_bytes):
    parsed = parse_document(sample_pdf_bytes, file_type="pdf", filename="annual.pdf")
    assert parsed.page_count == 2
    assert [p.page_number for p in parsed.pages] == [1, 2]
    assert "financial" in parsed.pages[0].text.lower()


def test_docx_parse_preserves_headings():
    from docx import Document as Docx

    doc = Docx()
    doc.add_heading("Company Policy", 0)
    doc.add_paragraph("All employees follow the data security guidelines.")
    doc.add_heading("Acceptable Use", 1)
    doc.add_paragraph("Work devices must stay locked.")
    buffer = io.BytesIO()
    doc.save(buffer)

    parsed = parse_document(buffer.getvalue(), file_type="docx", filename="policy.docx")
    assert parsed.file_type == "docx"
    assert "Company Policy" in parsed.full_text
    assert "Acceptable Use" in parsed.full_text
    assert "data security guidelines" in parsed.full_text


def test_unsupported_file_rejected():
    try:
        parse_document(b"not really", file_type="exe", filename="virus.exe")
        raise AssertionError("expected ValidationError")
    except Exception as exc:  # noqa: BLE001
        assert type(exc).__name__ == "ValidationError"


def test_extension_detection():
    assert document_parser.extension_of("report.PDF") == "pdf"
    assert document_parser.extension_of("notes.md") == "md"
    assert document_parser.extension_of("notes.markdown") == "md"
    assert document_parser.extension_of("contract.DOCX") == "docx"
    assert document_parser.extension_of("song.mp3") == ""
    assert document_parser.extension_of("noext") == ""
    assert not document_parser.is_supported_extension("a.bin")
    assert document_parser.is_supported_extension("a.pdf")
